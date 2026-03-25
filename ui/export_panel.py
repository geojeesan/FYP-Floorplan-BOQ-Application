import os
import traceback
import json
import pandas as pd
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QMessageBox, 
    QLabel, QFileDialog, QProgressBar, QComboBox
)
from PySide6.QtCore import Qt
import qtawesome as qta
from shapely.geometry import Point, Polygon as ShapelyPolygon

from .report_worker import ReportWorker
from database import get_item_details

try:
    import ezdxf
    from ezdxf.enums import TextEntityAlignment 
    EZDXF_AVAILABLE = True
except ImportError:
    EZDXF_AVAILABLE = False


class ExportPanel(QWidget):
    def __init__(self, main_window=None):
        super().__init__()
        self.main_window = main_window
        
        # Keep separate DataFrames for QTO and BOQ
        self.qto_df = None
        self.boq_df = None
        
        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignTop)
        layout.setSpacing(15)
        layout.setContentsMargins(5, 10, 5, 5)
        
        header_lbl = QLabel("<b>Export & Reports</b>")
        header_lbl.setStyleSheet("font-size: 16px; margin-bottom: 5px;")
        layout.addWidget(header_lbl)
        
        btn_style = """
            QPushButton {
                background-color: #fb9a44; color: white; font-size: 14px; 
                padding: 12px; border-radius: 6px; font-weight: bold;
                text-align: left; padding-left: 15px;
            }
            QPushButton:hover { background-color: #e08c3a; }
            QPushButton:disabled { background-color: #555555; color: #666666; }
        """
        cad_style = """
            QPushButton {
                text-align: left; font-size: 14px;
                padding: 12px; padding-left: 15px;
                color: white;
                font-weight: bold;
                border-radius: 6px;
            }
        """
        dl_csv_style = """
            QPushButton {
                text-align: center; 
                border: 1px solid rgba(96, 125, 139, 90); 
                padding: 6px;
                color: white;
                font-weight: bold;
                border-radius: 4px;
                margin-bottom: 25px;
            }
            QPushButton:hover { background-color: rgba(96, 125, 139, 90); }
        """

        dl_xl_style = """
            QPushButton {
                text-align: center; 
                border: 1px solid rgba(28, 136, 80, 90); 
                padding: 6px;
                color: white;
                font-weight: bold;
                border-radius: 4px;
                margin-bottom: 25px;
            }
            QPushButton:hover { background-color: rgba(28, 136, 80, 90); }
        """

        # QTO Section
        self.btn_qto = QPushButton(" Generate Quantity Take-off")
        self.btn_qto.setIcon(qta.icon('fa5s.file-excel', color='white'))
        self.btn_qto.setStyleSheet(btn_style)
        self.btn_qto.clicked.connect(lambda: self.start_hybrid_generation("Quantity Take-off (Measurement Sheet)"))
        
        self.qto_downloads_layout = QHBoxLayout()
        self.btn_qto_csv = QPushButton(" Download CSV")
        self.btn_qto_csv.setIcon(qta.icon('fa5s.file-csv', color='white'))
        self.btn_qto_csv.setStyleSheet(dl_csv_style)
        self.btn_qto_csv.clicked.connect(lambda: self.download_csv("QTO"))
        self.btn_qto_csv.hide()
        
        self.btn_qto_excel = QPushButton(" Download Excel")
        self.btn_qto_excel.setIcon(qta.icon('fa5s.file-excel', color='white'))
        self.btn_qto_excel.setStyleSheet(dl_xl_style)
        self.btn_qto_excel.clicked.connect(lambda: self.download_excel("QTO"))
        self.btn_qto_excel.hide()
        
        self.qto_downloads_layout.addWidget(self.btn_qto_csv)
        self.qto_downloads_layout.addWidget(self.btn_qto_excel)

        # BOQ Section
        boq_layout = QHBoxLayout()
        
        self.btn_boq = QPushButton(" Generate BOQ")
        self.btn_boq.setIcon(qta.icon('fa5s.file-invoice-dollar', color='white'))
        self.btn_boq.setStyleSheet(btn_style)
        self.btn_boq.clicked.connect(lambda: self.start_hybrid_generation("Bill of Quantities (BOQ)"))
        
        self.boq_standard_combo = QComboBox()
        # REMOVED CESMM4
        self.boq_standard_combo.addItems(["POMI", "NRM 2", "SMM7"]) 
        self.boq_standard_combo.setStyleSheet("""
            QComboBox {
                padding: 10px; border-radius: 6px; border: 0px solid #ccc; font-size: 14px;
            }
        """)
        
        boq_layout.addWidget(self.btn_boq, stretch=3)
        boq_layout.addWidget(self.boq_standard_combo, stretch=1)
        
        self.boq_downloads_layout = QHBoxLayout()
        self.btn_boq_csv = QPushButton(" Download CSV")
        self.btn_boq_csv.setIcon(qta.icon('fa5s.file-csv', color='white'))
        self.btn_boq_csv.setStyleSheet(dl_csv_style)
        self.btn_boq_csv.clicked.connect(lambda: self.download_csv("BOQ"))
        self.btn_boq_csv.hide()
        
        self.btn_boq_excel = QPushButton(" Download Excel")
        self.btn_boq_excel.setIcon(qta.icon('fa5s.file-excel', color='white'))
        self.btn_boq_excel.setStyleSheet(dl_xl_style)
        self.btn_boq_excel.clicked.connect(lambda: self.download_excel("BOQ"))
        self.btn_boq_excel.hide()
        
        self.boq_downloads_layout.addWidget(self.btn_boq_csv)
        self.boq_downloads_layout.addWidget(self.btn_boq_excel)

        # CAD Section
        self.btn_cad = QPushButton(" Export to CAD (.dxf)")
        self.btn_cad.setIcon(qta.icon('fa5s.drafting-compass', color='white'))
        self.btn_cad.setStyleSheet(cad_style)
        self.btn_cad.clicked.connect(self.export_cad)

        # Add everything to main layout
        layout.addWidget(self.btn_qto)
        layout.addLayout(self.qto_downloads_layout)
        
        layout.addLayout(boq_layout) 
        layout.addLayout(self.boq_downloads_layout)
        
        layout.addWidget(self.btn_cad)
        layout.addSpacing(20)
        
        # Progress UI
        self.progress_label = QLabel("")
        self.progress_label.setStyleSheet("color: #666; font-style: italic; font-size: 12px;")
        self.progress_label.setAlignment(Qt.AlignCenter)
        self.progress_label.hide()
        layout.addWidget(self.progress_label)

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setFixedHeight(8)
        self.progress_bar.setTextVisible(False)
        self.progress_bar.hide()
        layout.addWidget(self.progress_bar)

        layout.addStretch()

    def _prepare_hierarchical_data(self, viewer):
        """Builds ordered List: Room Header -> Materials -> Icons"""
        data = viewer.boq_data
        ratio = float(viewer.pixel_to_unit_ratio or 1.0)
        default_height = 2.4 

        room_totals = {}
        for item in data.get('rooms', []):
            label = item.get('label', 'Unknown Space')
            if label not in ["Wall", "Railing"]:
                room_totals[label] = room_totals.get(label, 0) + 1

        room_current = {}
        room_polys = []
        
        for item in data.get('rooms', []):
            label = item.get('label', 'Unknown Space')
            pts = item.get('points', [])
            
            if label not in ["Wall", "Railing"]:
                if room_totals.get(label, 0) > 1:
                    room_current[label] = room_current.get(label, 0) + 1
                    unique_label = f"{label} {room_current[label]}"
                else:
                    unique_label = label
                    
                if len(pts) >= 3:
                    try:
                        poly = ShapelyPolygon(pts)
                        room_polys.append((unique_label, poly, item))
                    except: pass
        
        unassigned_icons = []
        for item in data.get('icons', []):
            unassigned_icons.append({"label": item.get('label', 'Unknown Fixture'), "pts": item.get('points', []), "item": item, "matched": False})

        rows = []
        item_counter = 1

        for unique_label, poly, item in room_polys:
            pts = item.get('points', [])
            l, w, perimeter = "", "", 0
            if pts and len(pts) > 0:
                xs = [p[0] for p in pts]
                ys = [p[1] for p in pts]
                width_m = (max(xs) - min(xs)) * ratio
                length_m = (max(ys) - min(ys)) * ratio
                l = round(max(width_m, length_m), 2)
                w = round(min(width_m, length_m), 2)
                perimeter = round(poly.length * ratio, 2)

            area_px = item.get('area_pixels', 0)
            area_val = round(area_px * (ratio ** 2), 2)
            
            rows.append({
                "Type": "Header", "Heading": unique_label, "Base_Description": "", "Category": "Space/Room",
                "Unit": "sqm", "Length": l, "Width": w, "Height": default_height, "Quantity": area_val, 
                "Rate": "", "Total_Amount": "", "Item_No": ""
            })
            
            materials = item.get('materials', {})
            has_flooring = False
            has_wall = False
            
            for cat, mat_data in materials.items():
                desc = mat_data.get('name', '')
                
                db_record = get_item_details(desc)
                if not db_record: db_record = {} 

                brand = db_record.get('Brand_Name', '')
                if brand and brand.lower() not in desc.lower():
                    desc = f"{brand} {desc}"
                
                official_item_no = db_record.get('Item_No', f"EST-{item_counter:03d}")
                raw_unit = db_record.get('Unit', mat_data.get('unit', 'sqm')).lower()
                markup = float(db_record.get('Markup_Percentage', 0.0))
                rate = float(db_record.get('Cost_per_Unit', mat_data.get('cost', 0.0)))
                
                if raw_unit in ['litre', 'litres', 'l']:
                    unit = 'litre'
                    floor_qty = round(area_val / 5.0, 2)
                    wall_qty = round((perimeter * default_height) / 5.0, 2)
                else:
                    unit = db_record.get('Unit', 'sqm')
                    floor_qty = area_val
                    wall_qty = round(perimeter * default_height, 2)

                if "Floor" in cat or "floor" in cat.lower():
                    has_flooring = True
                    total_amt = round(rate * floor_qty, 2)
                    final_amt = round(total_amt + (total_amt * (markup / 100.0)), 2)

                    rows.append({
                        "Type": "Sub", "Heading": "", "Base_Description": desc, "Category": cat,
                        "Unit": unit, "Length": l, "Width": w, "Height": "", "Quantity": floor_qty,
                        "Rate": rate, "Total_Amount": total_amt, 
                        "Markup_Percentage": markup, "Final_Amount": final_amt, "Item_No": official_item_no
                    })
                    item_counter += 1
                elif "Wall" in cat or "wall" in cat.lower():
                    has_wall = True
                    total_amt = round(rate * wall_qty, 2)
                    final_amt = round(total_amt + (total_amt * (markup / 100.0)), 2)

                    rows.append({
                        "Type": "Sub", "Heading": "", "Base_Description": desc, "Category": cat,
                        "Unit": unit, "Length": perimeter, "Width": "", "Height": default_height, "Quantity": wall_qty,
                        "Rate": rate, "Total_Amount": total_amt, 
                        "Markup_Percentage": markup, "Final_Amount": final_amt, "Item_No": official_item_no
                    })
                    item_counter += 1

            if not has_flooring:
                rows.append({
                    "Type": "Sub", "Heading": "", "Base_Description": "Flooring Finish", "Category": "Flooring",
                    "Unit": "sqm", "Length": l, "Width": w, "Height": "", "Quantity": area_val,
                    "Rate": 0.0, "Total_Amount": 0.0, "Item_No": f"EST-{item_counter:03d}"
                })
                item_counter += 1
                
            if not has_wall:
                wall_area = round(perimeter * default_height, 2)
                rows.append({
                    "Type": "Sub", "Heading": "", "Base_Description": "Wall Finish", "Category": "Wall Finish",
                    "Unit": "sqm", "Length": perimeter, "Width": "", "Height": default_height, "Quantity": wall_area,
                    "Rate": 0.0, "Total_Amount": 0.0, "Item_No": f"EST-{item_counter:03d}"
                })
                item_counter += 1

            for icon_obj in unassigned_icons:
                if not icon_obj["matched"] and len(icon_obj["pts"]) > 0:
                    pts = icon_obj["pts"]
                    cx = sum([p[0] for p in pts]) / len(pts)
                    cy = sum([p[1] for p in pts]) / len(pts)
                    pt = Point(cx, cy)
                    
                    if poly.contains(pt) or poly.distance(pt) < 30:
                        icon_obj["matched"] = True
                        label = icon_obj["label"]
                        icon_item = icon_obj["item"]
                        i_mats = icon_item.get('materials', {})
                        
                        icon_width = ""
                        if "door" in label.lower() or "window" in label.lower():
                            xs = [p[0] for p in pts]
                            ys = [p[1] for p in pts]
                            w_m = (max(xs) - min(xs)) * ratio
                            h_m = (max(ys) - min(ys)) * ratio
                            icon_width = round(max(w_m, h_m), 2)
                        
                        if i_mats:
                            for c, m in i_mats.items():
                                mat_name = m.get('name', '')
                                desc = f"{label} - {mat_name}"
                                db_record = get_item_details(mat_name) or get_item_details(label) or {}

                                brand = db_record.get('Brand_Name', '')
                                if brand and brand.lower() not in desc.lower():
                                    desc = f"{brand} {desc}"
                                
                                official_item_no = db_record.get('Item_No', f"EST-{item_counter:03d}")
                                unit = db_record.get('Unit', m.get('unit', 'ea'))
                                markup = float(db_record.get('Markup_Percentage', 0.0))
                                rate = float(db_record.get('Cost_per_Unit', m.get('cost', 0.0)))
                                
                                total_amt = round(rate * 1.0, 2)
                                final_amt = round(total_amt + (total_amt * (markup / 100.0)), 2)
                                
                                rows.append({
                                    "Type": "Sub", "Heading": "", "Base_Description": desc, "Category": c,
                                    "Unit": unit, "Length": "", "Width": icon_width, "Height": "", "Quantity": 1.0,
                                    "Rate": rate, "Total_Amount": total_amt, 
                                    "Markup_Percentage": markup, "Final_Amount": final_amt, "Item_No": official_item_no
                                })
                                item_counter += 1
                        else:
                            official_item_no = f"EST-{item_counter:03d}"
                            unit = 'ea'
                            markup = 0.0
                            rate = 0.0
                            total_amt = 0.0
                            final_amt = 0.0

                            rows.append({
                                "Type": "Sub", "Heading": "", "Base_Description": label, "Category": "Fixture/Item",
                                "Unit": unit, "Length": "", "Width": icon_width, "Height": "", "Quantity": 1.0,
                                "Rate": rate, "Total_Amount": total_amt, 
                                "Markup_Percentage": markup, "Final_Amount": final_amt, "Item_No": official_item_no
                            })
                            item_counter += 1

        return rows

    def start_hybrid_generation(self, report_type):
        viewer = self.main_window.current_widget() if self.main_window else None
        
        if not viewer or not getattr(viewer, 'boq_data', None):
            QMessageBox.warning(self, "Export Error", "No active analysis data to export.")
            return
            
        self.btn_qto.setEnabled(False)
        self.btn_boq.setEnabled(False)
        self.progress_label.setText("Starting process...")
        self.progress_label.show()
        self.progress_bar.show()
        
        provider = self.main_window.chat_panel.provider_combo.currentText()
        model = self.main_window.chat_panel.model_combo.currentText()
        boq_standard = self.boq_standard_combo.currentText()
        
        exact_math_list = self._prepare_hierarchical_data(viewer)
            
        self.worker = ReportWorker(exact_math_list, report_type, provider, model, boq_standard)
        self.worker.progress.connect(self.update_progress)
        self.worker.finished.connect(self.on_report_finished)
        self.worker.error.connect(self.on_report_error)
        self.worker.start()

    def update_progress(self, msg):
        self.progress_label.setText(msg)

    def on_report_finished(self, json_str, report_type):
        self.progress_bar.hide()
        self.progress_label.hide()
        self.btn_qto.setEnabled(True)
        self.btn_boq.setEnabled(True)
        
        try:
            data = json.loads(json_str)
            df = pd.DataFrame(data)
            
            if "Take-off" in report_type:
                self.qto_df = df
                self.btn_qto_csv.show()
                self.btn_qto_excel.show()
            else:
                self.boq_df = df
                self.btn_boq_csv.show()
                self.btn_boq_excel.show()
                
            QMessageBox.information(self, "Success", f"{report_type} generated successfully!\n\nPlease choose a download format")
            
        except Exception as e:
            QMessageBox.critical(self, "AI Formatting Error", f"The AI failed to format the report properly as JSON.\nError: {e}\n\nRaw Output:\n{json_str[:500]}...")

    def on_report_error(self, err_msg):
        self.progress_bar.hide()
        self.progress_label.hide()
        self.btn_qto.setEnabled(True)
        self.btn_boq.setEnabled(True)
        QMessageBox.critical(self, "Generation Error", f"An error occurred during report generation:\n{err_msg}")

    def download_csv(self, target_type):
        df = self.qto_df if target_type == "QTO" else self.boq_df
        if df is None: return
        
        file_path, _ = QFileDialog.getSaveFileName(self, "Save CSV", f"{target_type}_Report.csv", "CSV Files (*.csv)")
        if file_path:
            try:
                df.to_csv(file_path, index=False)
                QMessageBox.information(self, "Saved", f"{target_type} CSV File saved successfully!")
            except Exception as e:
                QMessageBox.critical(self, "Error", f"Failed to save CSV:\n{str(e)}")

    def download_excel(self, target_type):
        df = self.qto_df if target_type == "QTO" else self.boq_df
        if df is None: return
        
        file_path, _ = QFileDialog.getSaveFileName(self, "Save Excel", f"{target_type}_Report.xlsx", "Excel Files (*.xlsx)")
        if file_path:
            try:
                df.to_excel(file_path, index=False)
                
                try:
                    from openpyxl import load_workbook
                    from openpyxl.styles import Font
                    
                    wb = load_workbook(file_path)
                    ws = wb.active
                    
                    for cell in ws[1]:
                        cell.font = Font(bold=True)
                        
                    heading_col_idx = None
                    for col_idx, cell in enumerate(ws[1], 1):
                        if cell.value == "Heading":
                            heading_col_idx = col_idx
                            break
                            
                    if heading_col_idx:
                        for row in ws.iter_rows(min_row=2):
                            heading_cell = row[heading_col_idx - 1]
                            if heading_cell.value and str(heading_cell.value).strip():
                                for cell in row:
                                    cell.font = Font(bold=True)
                                    
                    wb.save(file_path)
                except ImportError:
                    pass
                except Exception as e:
                    print(f"Failed to apply Excel styling: {e}")
                
                QMessageBox.information(self, "Saved", f"{target_type} Excel File saved successfully!")
            except Exception as e:
                QMessageBox.critical(self, "Error", f"Failed to save Excel:\n{str(e)}")

    def export_cad(self):
        if not EZDXF_AVAILABLE:
            QMessageBox.warning(self, "Missing Dependency", "The 'ezdxf' library is required to export to CAD.\n\nPlease install it using your terminal:\npip install ezdxf")
            return
            
        viewer = self.main_window.current_widget() if self.main_window else None
        
        if not viewer or not getattr(viewer, 'boq_data', None):
            QMessageBox.warning(self, "Export Error", "No active analysis data to export.")
            return
            
        boq_data = viewer.boq_data
        ratio = float(getattr(viewer, 'pixel_to_unit_ratio', 1.0) or 1.0)
            
        file_path, _ = QFileDialog.getSaveFileName(self, "Save CAD File", "", "DXF Files (*.dxf)")
        if not file_path: return
            
        try:
            doc = ezdxf.new('R2010')
            msp = doc.modelspace()
            
            doc.layers.add(name="Rooms", color=3)       
            doc.layers.add(name="Structures", color=1)  
            doc.layers.add(name="Icons", color=5)       
            doc.layers.add(name="Labels", color=7)      
            
            text_height = float(15 * ratio)
            
            def add_polygon_to_msp(item_list, default_layer):
                for item in item_list:
                    points = item.get('points', [])
                    if not points or len(points) < 3: continue
                        
                    cad_points = [(float(p[0] * ratio), float(-p[1] * ratio)) for p in points]
                    label = str(item.get('label', 'Unknown'))
                    layer = default_layer
                    if default_layer == "Rooms" and label in ["Wall", "Railing"]:
                        layer = "Structures"
                        
                    msp.add_lwpolyline(cad_points, close=True, dxfattribs={'layer': layer})
                    cx = float(sum([p[0] for p in cad_points]) / len(cad_points))
                    cy = float(sum([p[1] for p in cad_points]) / len(cad_points))
                    
                    msp.add_text(
                        label, 
                        dxfattribs={'layer': 'Labels', 'height': text_height}
                    ).set_placement((cx, cy), align=TextEntityAlignment.MIDDLE_CENTER)
                    
            add_polygon_to_msp(boq_data.get('rooms', []), "Rooms")
            add_polygon_to_msp(boq_data.get('icons', []), "Icons")
            
            if hasattr(viewer, 'manual_text_labels'):
                for txt_item in viewer.manual_text_labels:
                    x, y = float(txt_item['pos'].x() * ratio), float(-txt_item['pos'].y() * ratio)
                    msp.add_text(
                        str(txt_item['text']), 
                        dxfattribs={'layer': 'Labels', 'height': text_height}
                    ).set_placement((x, y), align=TextEntityAlignment.MIDDLE_CENTER)

            doc.saveas(file_path)
            QMessageBox.information(self, "Export Success", f"Successfully exported CAD file to:\n{file_path}")
            
        except Exception as e:
            error_trace = traceback.format_exc()
            print("CAD Export Failed:\n", error_trace)
            QMessageBox.critical(
                self, 
                "Export Error", 
                f"Failed to generate DXF file:\n{str(e)}\n\nCheck your terminal for the full traceback details."
            )